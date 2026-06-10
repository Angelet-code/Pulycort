from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    KeepTogether,
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "04_analisis_y_entregables" / "informes"
DOCX_PATH = OUT_DIR / "solicitud_documentacion_mapeo_empresa_pulycort.docx"
PDF_PATH = OUT_DIR / "solicitud_documentacion_mapeo_empresa_pulycort.pdf"

BLUE = "1F4D78"
MID_BLUE = "2E74B5"
DARK = "172033"
MUTED = "5A6678"
LINE = "D7DEE8"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
SOFT_GOLD = "FFF4DD"
GOLD = "7A5A00"
WHITE = "FFFFFF"
FONT = "Arial"
PDF_FONT = "ArialEmbed"
PDF_FONT_BOLD = "ArialEmbed-Bold"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = LINE, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(indent_dxa))

    old_grid = tbl.find(qn("w:tblGrid"))
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    tbl.insert(0, grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            set_cell_margins(cell)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_paragraph_style(paragraph, size=10.5, color=DARK, bold=False, before=0, after=6, line=1.1):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    for run in paragraph.runs:
        run.font.name = FONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)
        run.bold = bold


def add_run(paragraph, text: str, *, bold=False, color=DARK, size=10.5):
    run = paragraph.add_run(text)
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    return run


def add_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_paragraph()
    if level == 1:
        add_run(paragraph, text, bold=True, color=MID_BLUE, size=16)
        paragraph.paragraph_format.space_before = Pt(16)
        paragraph.paragraph_format.space_after = Pt(8)
    elif level == 2:
        add_run(paragraph, text, bold=True, color=MID_BLUE, size=13)
        paragraph.paragraph_format.space_before = Pt(12)
        paragraph.paragraph_format.space_after = Pt(6)
    else:
        add_run(paragraph, text, bold=True, color=BLUE, size=12)
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None):
    paragraph = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        add_run(paragraph, bold_prefix, bold=True)
        add_run(paragraph, text[len(bold_prefix):])
    else:
        add_run(paragraph, text)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    return paragraph


def add_bullet(doc: Document, text: str, *, style="List Bullet"):
    paragraph = doc.add_paragraph(style=style)
    add_run(paragraph, text)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.15
    return paragraph


def style_table_header(row) -> None:
    for cell in row.cells:
        set_cell_shading(cell, LIGHT_GRAY)
        for paragraph in cell.paragraphs:
            if not paragraph.runs:
                continue
            set_paragraph_style(paragraph, size=9.5, color=BLUE, bold=True, after=0, line=1.05)


def style_table_body(table, priority_col: int | None = None) -> None:
    for r_idx, row in enumerate(table.rows):
        if r_idx == 0:
            continue
        for c_idx, cell in enumerate(row.cells):
            if priority_col is not None and c_idx == priority_col:
                set_cell_shading(cell, SOFT_GOLD)
            for paragraph in cell.paragraphs:
                if not paragraph.runs:
                    continue
                color = GOLD if priority_col is not None and c_idx == priority_col else DARK
                bold = priority_col is not None and c_idx == priority_col
                set_paragraph_style(paragraph, size=9, color=color, bold=bold, after=0, line=1.08)


def add_callout(doc: Document, title: str, body: str, fill: str = LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360], indent_dxa=120)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    add_run(paragraph, title + " ", bold=True, color=BLUE, size=10.5)
    add_run(paragraph, body, color=DARK, size=10.5)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.12
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.15

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(header, "Pulycort | Solicitud de documentación", color=MUTED, size=8.5)
    header.paragraph_format.space_after = Pt(0)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(footer, "Documento de trabajo para mapeo operativo y preparación de integraciones", color=MUTED, size=8)
    footer.paragraph_format.space_after = Pt(0)


def build_document() -> Document:
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    add_run(title, "Solicitud de documentación para el mapeo operativo de Pulycort", bold=True, color=BLUE, size=22)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.line_spacing = 1.0

    subtitle = doc.add_paragraph()
    add_run(
        subtitle,
        "Objetivo: comprender procesos, trazabilidad, producción, Odoo, máquinas y gestión diaria para preparar mejoras operativas y futuras integraciones.",
        color=MUTED,
        size=11.5,
    )
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle.paragraph_format.line_spacing = 1.15

    meta = doc.add_table(rows=3, cols=2)
    set_table_geometry(meta, [1800, 7560], indent_dxa=120)
    meta_data = [
        ("Dirigido a", "Equipo de Pulycort / INDASEL"),
        ("Fecha", "10 de junio de 2026"),
        ("Finalidad", "Reunir ejemplos reales y documentación completa para mapear la empresa de punta a punta."),
    ]
    for row, (key, value) in zip(meta.rows, meta_data):
        row.cells[0].text = key
        row.cells[1].text = value
        set_cell_shading(row.cells[0], LIGHT_GRAY)
        for paragraph in row.cells[0].paragraphs:
            set_paragraph_style(paragraph, size=9, color=BLUE, bold=True, after=0)
        for paragraph in row.cells[1].paragraphs:
            set_paragraph_style(paragraph, size=9, color=DARK, after=0)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_body(
        doc,
        "Para avanzar con rigor en el mapeo operativo de Pulycort, necesitamos trabajar con documentación real de la empresa: expedientes, partes, capturas, exportaciones y ejemplos de casos habituales y excepcionales.",
    )
    add_body(
        doc,
        "No es necesario que la información esté perfecta ni homogeneizada. También son útiles documentos incompletos, capturas de pantalla, hojas internas, archivos desordenados o ejemplos donde se vea cómo se resuelven los problemas en la práctica.",
    )
    add_callout(
        doc,
        "Envío y seguridad:",
        "solicitamos documentación completa. Los archivos con datos personales, clientes, precios, accesos, IPs, credenciales o información sensible deben compartirse únicamente por el canal seguro que acordemos.",
        fill=SOFT_GOLD,
    )

    add_heading(doc, "1. Documentación prioritaria", 1)
    add_body(
        doc,
        "Estos bloques son los más importantes para pasar de entender catálogos a entender cómo fluye realmente el negocio.",
    )

    priority_table = doc.add_table(rows=1, cols=3)
    priority_table.rows[0].cells[0].text = "Documentación solicitada"
    priority_table.rows[0].cells[1].text = "Para qué sirve"
    priority_table.rows[0].cells[2].text = "Prioridad"
    priority_rows = [
        (
            "5-10 pedidos reales completos de punta a punta",
            "Ver el recorrido completo desde presupuesto hasta cobro: pedido, producción, lote/PM, albarán, factura, incidencias y comunicaciones relevantes.",
            "Muy alta",
        ),
        (
            "Partes reales y exportaciones de producción por máquina",
            "Entender campos, tiempos, operarios, materiales, lotes, consumos, estados e identificadores que conectan producción con Odoo.",
            "Muy alta",
        ),
        (
            "Capturas o exportaciones del sistema actual",
            "Ver cómo están modelados hoy productos, variantes, lotes, stock, pedidos, clientes, tarifas, partes y máquinas.",
            "Alta",
        ),
    ]
    for item, purpose, priority in priority_rows:
        cells = priority_table.add_row().cells
        cells[0].text = item
        cells[1].text = purpose
        cells[2].text = priority
    set_table_geometry(priority_table, [2950, 5030, 1380], indent_dxa=120)
    style_table_header(priority_table.rows[0])
    style_table_body(priority_table, priority_col=2)

    add_heading(doc, "2. Documentación operativa", 1)
    operational_table = doc.add_table(rows=1, cols=3)
    operational_table.rows[0].cells[0].text = "Bloque"
    operational_table.rows[0].cells[1].text = "Qué pedir"
    operational_table.rows[0].cells[2].text = "Objetivo"
    operational_rows = [
        (
            "Organización y responsabilidades",
            "Roles, responsables por área, quién presupuesta, quién confirma pedidos, quién planifica, quién produce, quién valida calidad y quién factura.",
            "Entender propietarios, decisiones y puntos de coordinación.",
        ),
        (
            "Planificación y control diario",
            "Calendarios de producción, hojas de seguimiento, pizarras, Excel internos, listas de tareas, partes manuales y comunicaciones recurrentes.",
            "Detectar los sistemas reales que sostienen la operación diaria.",
        ),
        (
            "Casos problemáticos o excepciones",
            "Pedidos con medidas especiales, cambios de material, roturas, mermas, reprocesos, retrasos, errores de stock o cambios de pedido.",
            "Diseñar procesos que aguanten la realidad, no solo el caso ideal.",
        ),
        (
            "Inventario físico y trazabilidad",
            "Etiquetas, códigos de bloque, tabla, losa, palet, cajón, ubicación, fotos de identificación y ejemplos de movimientos.",
            "Conectar material físico, lote, ubicación y pedido.",
        ),
    ]
    for block, request, objective in operational_rows:
        cells = operational_table.add_row().cells
        cells[0].text = block
        cells[1].text = request
        cells[2].text = objective
    set_table_geometry(operational_table, [2300, 4300, 2760], indent_dxa=120)
    style_table_header(operational_table.rows[0])
    style_table_body(operational_table)

    add_heading(doc, "3. Información comercial, económica y de precio", 1)
    add_body(
        doc,
        "Para que el modelo operativo sea útil, también necesitamos comprender cómo se decide y se calcula el valor económico de cada trabajo.",
    )
    add_bullet(doc, "Reglas reales de presupuesto: qué datos se usan, qué se calcula automáticamente y qué se decide manualmente.")
    add_bullet(doc, "Tarifas aplicadas por material, acabado, grosor, medida, familia, cliente, urgencia, transporte o embalaje.")
    add_bullet(doc, "Descuentos, condiciones especiales, excepciones habituales y criterios para modificar un precio estándar.")
    add_bullet(doc, "Costes directos e indirectos que se consideren relevantes para producción, máquina, operación o margen.")

    add_heading(doc, "4. Formato de entrega recomendado", 1)
    add_body(
        doc,
        "Lo ideal es conservar los nombres originales de los archivos y acompañarlos de una breve nota cuando el contexto no sea evidente. Si existen exportaciones, preferimos recibir el formato original además de PDF o captura.",
    )
    format_table = doc.add_table(rows=1, cols=2)
    format_table.rows[0].cells[0].text = "Tipo de material"
    format_table.rows[0].cells[1].text = "Formato útil"
    format_rows = [
        ("Expedientes de pedido", "Carpeta por pedido con presupuesto, pedido, producción, albarán, factura y comunicaciones asociadas."),
        ("Partes de máquina", "Archivo original exportado, Excel/CSV si existe, PDF o capturas si no hay exportación."),
        ("Sistemas actuales", "Capturas de pantalla con contexto: módulo, vista, filtros usados y ejemplo concreto."),
        ("Documentos internos", "Excel, Word, PDF, fotos o capturas tal como se usan actualmente."),
    ]
    for material_type, useful_format in format_rows:
        cells = format_table.add_row().cells
        cells[0].text = material_type
        cells[1].text = useful_format
    set_table_geometry(format_table, [2500, 6860], indent_dxa=120)
    style_table_header(format_table.rows[0])
    style_table_body(format_table)

    add_heading(doc, "5. Primer envío recomendado", 1)
    add_body(
        doc,
        "Para empezar sin bloquear el trabajo, proponemos un primer paquete acotado con los siguientes elementos:",
    )
    first_pack = [
        "5-10 pedidos reales completos, incluyendo algún caso sencillo y algún caso problemático.",
        "Exportaciones o partes reales de telar, reforzadora, pulidora, disco puente y taller, si están disponibles.",
        "Capturas de las pantallas principales del sistema actual: pedido, producto, lote, stock, parte de producción, cliente y factura.",
        "Ejemplos de etiquetas físicas y trazabilidad de bloque, tabla, losa, palet, cajón y ubicación.",
        "Una breve lista de responsables por área y de las decisiones que toma cada uno.",
    ]
    for item in first_pack:
        add_bullet(doc, item)

    add_callout(
        doc,
        "Criterio práctico:",
        "es preferible recibir documentación real aunque esté incompleta antes que esperar a preparar un paquete perfecto. Los huecos, contradicciones y formatos manuales también ayudan a detectar dónde puede aportar valor la automatización.",
    )

    add_heading(doc, "Cierre", 1)
    add_body(
        doc,
        "Con esta documentación podremos construir un mapa operativo fiable de Pulycort, validar el modelo de datos para Odoo, identificar puntos críticos de trazabilidad y priorizar pequeñas soluciones prácticas antes de abordar integraciones mayores.",
    )
    add_body(
        doc,
        "El primer objetivo no es auditar ni juzgar la forma actual de trabajo, sino entenderla con precisión para diseñar mejoras que encajen con la realidad de la empresa.",
    )

    return doc


def register_pdf_fonts() -> tuple[str, str]:
    regular = Path("C:/Windows/Fonts/arial.ttf")
    bold = Path("C:/Windows/Fonts/arialbd.ttf")
    if regular.exists() and bold.exists():
        pdfmetrics.registerFont(TTFont(PDF_FONT, str(regular)))
        pdfmetrics.registerFont(TTFont(PDF_FONT_BOLD, str(bold)))
        return PDF_FONT, PDF_FONT_BOLD
    return "Helvetica", "Helvetica-Bold"


def hc(value: str):
    return colors.HexColor("#" + value)


def pdf_paragraph(text: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(text.replace("\n", "<br/>"), style)


def pdf_table(rows: list[list[str]], widths: list[float], styles, header_fill=LIGHT_GRAY, priority_col: int | None = None):
    data = []
    header_style = styles["TableHeader"]
    cell_style = styles["TableCell"]
    priority_style = styles["PriorityCell"]
    for row_idx, row in enumerate(rows):
        data_row = []
        for col_idx, value in enumerate(row):
            style = header_style if row_idx == 0 else priority_style if priority_col == col_idx else cell_style
            data_row.append(pdf_paragraph(value, style))
        data.append(data_row)
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), hc(header_fill)),
        ("TEXTCOLOR", (0, 0), (-1, 0), hc(BLUE)),
        ("GRID", (0, 0), (-1, -1), 0.45, hc(LINE)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
    ]
    if priority_col is not None:
        commands.append(("BACKGROUND", (priority_col, 1), (priority_col, -1), hc(SOFT_GOLD)))
    table.setStyle(TableStyle(commands))
    return table


def add_pdf_callout(story, title: str, body: str, styles, fill=LIGHT_BLUE) -> None:
    text = f"<b>{title}</b> {body}"
    table = Table(
        [[pdf_paragraph(text, styles["Callout"])]],
        colWidths=[17.4 * cm],
        hAlign="LEFT",
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), hc(fill)),
                ("BOX", (0, 0), (-1, -1), 0.45, hc(LINE)),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ("LEFTPADDING", (0, 0), (-1, -1), 9),
                ("RIGHTPADDING", (0, 0), (-1, -1), 9),
            ]
        )
    )
    story.extend([table, Spacer(1, 8)])


def build_pdf() -> None:
    font, font_bold = register_pdf_fonts()
    base = getSampleStyleSheet()
    styles = {
        "Title": ParagraphStyle(
            "Title",
            parent=base["Title"],
            fontName=font_bold,
            fontSize=24,
            leading=26,
            textColor=hc(BLUE),
            alignment=TA_LEFT,
            spaceAfter=10,
        ),
        "Subtitle": ParagraphStyle(
            "Subtitle",
            parent=base["BodyText"],
            fontName=font,
            fontSize=11.5,
            leading=14,
            textColor=hc(MUTED),
            spaceAfter=12,
        ),
        "Body": ParagraphStyle(
            "Body",
            parent=base["BodyText"],
            fontName=font,
            fontSize=10.4,
            leading=13.0,
            textColor=hc(DARK),
            spaceAfter=6,
        ),
        "H1": ParagraphStyle(
            "H1",
            parent=base["Heading1"],
            fontName=font_bold,
            fontSize=16,
            leading=19,
            textColor=hc(MID_BLUE),
            spaceBefore=15,
            spaceAfter=7,
            keepWithNext=True,
        ),
        "H2": ParagraphStyle(
            "H2",
            parent=base["Heading2"],
            fontName=font_bold,
            fontSize=12.5,
            leading=15,
            textColor=hc(MID_BLUE),
            spaceBefore=10,
            spaceAfter=5,
            keepWithNext=True,
        ),
        "TableHeader": ParagraphStyle(
            "TableHeader",
            parent=base["BodyText"],
            fontName=font_bold,
            fontSize=8.8,
            leading=10.5,
            textColor=hc(BLUE),
            spaceAfter=0,
        ),
        "TableCell": ParagraphStyle(
            "TableCell",
            parent=base["BodyText"],
            fontName=font,
            fontSize=8.7,
            leading=10.5,
            textColor=hc(DARK),
            spaceAfter=0,
        ),
        "PriorityCell": ParagraphStyle(
            "PriorityCell",
            parent=base["BodyText"],
            fontName=font_bold,
            fontSize=8.7,
            leading=10.5,
            textColor=hc(GOLD),
            alignment=TA_CENTER,
            spaceAfter=0,
        ),
        "Callout": ParagraphStyle(
            "Callout",
            parent=base["BodyText"],
            fontName=font,
            fontSize=10,
            leading=12.5,
            textColor=hc(DARK),
            spaceAfter=0,
        ),
        "Bullet": ParagraphStyle(
            "Bullet",
            parent=base["BodyText"],
            fontName=font,
            fontSize=10,
            leading=12.5,
            leftIndent=15,
            bulletIndent=4,
            textColor=hc(DARK),
            spaceAfter=5,
        ),
        "MetaKey": ParagraphStyle(
            "MetaKey",
            parent=base["BodyText"],
            fontName=font_bold,
            fontSize=8.8,
            leading=10.5,
            textColor=hc(BLUE),
            spaceAfter=0,
        ),
        "MetaVal": ParagraphStyle(
            "MetaVal",
            parent=base["BodyText"],
            fontName=font,
            fontSize=8.8,
            leading=10.5,
            textColor=hc(DARK),
            spaceAfter=0,
        ),
    }

    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=A4,
        rightMargin=1.6 * cm,
        leftMargin=1.6 * cm,
        topMargin=1.55 * cm,
        bottomMargin=1.45 * cm,
        title="Solicitud de documentación para el mapeo operativo de Pulycort",
        author="Pulycort",
    )

    story = [
        pdf_paragraph("Solicitud de documentación para el mapeo operativo de Pulycort", styles["Title"]),
        pdf_paragraph(
            "Objetivo: comprender procesos, trazabilidad, producción, Odoo, máquinas y gestión diaria para preparar mejoras operativas y futuras integraciones.",
            styles["Subtitle"],
        ),
    ]

    meta = [
        ["Dirigido a", "Equipo de Pulycort / INDASEL"],
        ["Fecha", "10 de junio de 2026"],
        ["Finalidad", "Reunir ejemplos reales y documentación completa para mapear la empresa de punta a punta."],
    ]
    meta_data = [
        [pdf_paragraph(key, styles["MetaKey"]), pdf_paragraph(value, styles["MetaVal"])]
        for key, value in meta
    ]
    meta_table = Table(meta_data, colWidths=[3.0 * cm, 14.4 * cm], hAlign="LEFT")
    meta_table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.45, hc(LINE)),
                ("BACKGROUND", (0, 0), (0, -1), hc(LIGHT_GRAY)),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ]
        )
    )
    story.extend([meta_table, Spacer(1, 12)])

    story.extend(
        [
            pdf_paragraph(
                "Para avanzar con rigor en el mapeo operativo de Pulycort, necesitamos trabajar con documentación real de la empresa: expedientes, partes, capturas, exportaciones y ejemplos de casos habituales y excepcionales.",
                styles["Body"],
            ),
            pdf_paragraph(
                "No es necesario que la información esté perfecta ni homogeneizada. También son útiles documentos incompletos, capturas de pantalla, hojas internas, archivos desordenados o ejemplos donde se vea cómo se resuelven los problemas en la práctica.",
                styles["Body"],
            ),
        ]
    )
    add_pdf_callout(
        story,
        "Envío y seguridad:",
        "solicitamos documentación completa. Los archivos con datos personales, clientes, precios, accesos, IPs, credenciales o información sensible deben compartirse únicamente por el canal seguro que acordemos.",
        styles,
        fill=SOFT_GOLD,
    )

    story.append(pdf_paragraph("1. Documentación prioritaria", styles["H1"]))
    story.append(
        pdf_paragraph(
            "Estos bloques son los más importantes para pasar de entender catálogos a entender cómo fluye realmente el negocio.",
            styles["Body"],
        )
    )
    story.append(
        pdf_table(
            [
                ["Documentación solicitada", "Para qué sirve", "Prioridad"],
                [
                    "5-10 pedidos reales completos de punta a punta",
                    "Ver el recorrido completo desde presupuesto hasta cobro: pedido, producción, lote/PM, albarán, factura, incidencias y comunicaciones relevantes.",
                    "Muy alta",
                ],
                [
                    "Partes reales y exportaciones de producción por máquina",
                    "Entender campos, tiempos, operarios, materiales, lotes, consumos, estados e identificadores que conectan producción con Odoo.",
                    "Muy alta",
                ],
                [
                    "Capturas o exportaciones del sistema actual",
                    "Ver cómo están modelados hoy productos, variantes, lotes, stock, pedidos, clientes, tarifas, partes y máquinas.",
                    "Alta",
                ],
            ],
            [5.1 * cm, 9.9 * cm, 2.4 * cm],
            styles,
            priority_col=2,
        )
    )

    story.append(pdf_paragraph("2. Documentación operativa", styles["H1"]))
    story.append(
        pdf_table(
            [
                ["Bloque", "Qué pedir", "Objetivo"],
                [
                    "Organización y responsabilidades",
                    "Roles, responsables por área, quién presupuesta, quién confirma pedidos, quién planifica, quién produce, quién valida calidad y quién factura.",
                    "Entender propietarios, decisiones y puntos de coordinación.",
                ],
                [
                    "Planificación y control diario",
                    "Calendarios de producción, hojas de seguimiento, pizarras, Excel internos, listas de tareas, partes manuales y comunicaciones recurrentes.",
                    "Detectar los sistemas reales que sostienen la operación diaria.",
                ],
                [
                    "Casos problemáticos o excepciones",
                    "Pedidos con medidas especiales, cambios de material, roturas, mermas, reprocesos, retrasos, errores de stock o cambios de pedido.",
                    "Diseñar procesos que aguanten la realidad, no solo el caso ideal.",
                ],
                [
                    "Inventario físico y trazabilidad",
                    "Etiquetas, códigos de bloque, tabla, losa, palet, cajón, ubicación, fotos de identificación y ejemplos de movimientos.",
                    "Conectar material físico, lote, ubicación y pedido.",
                ],
            ],
            [4.1 * cm, 8.0 * cm, 5.3 * cm],
            styles,
        )
    )

    story.append(pdf_paragraph("3. Información comercial, económica y de precio", styles["H1"]))
    story.append(
        pdf_paragraph(
            "Para que el modelo operativo sea útil, también necesitamos comprender cómo se decide y se calcula el valor económico de cada trabajo.",
            styles["Body"],
        )
    )
    commercial_items = [
        "Reglas reales de presupuesto: qué datos se usan, qué se calcula automáticamente y qué se decide manualmente.",
        "Tarifas aplicadas por material, acabado, grosor, medida, familia, cliente, urgencia, transporte o embalaje.",
        "Descuentos, condiciones especiales, excepciones habituales y criterios para modificar un precio estándar.",
        "Costes directos e indirectos que se consideren relevantes para producción, máquina, operación o margen.",
    ]
    story.append(
        ListFlowable(
            [ListItem(pdf_paragraph(item, styles["Bullet"])) for item in commercial_items],
            bulletType="bullet",
            start="circle",
            leftIndent=12,
            bulletFontName=font,
        )
    )

    story.append(pdf_paragraph("4. Formato de entrega recomendado", styles["H1"]))
    story.append(
        pdf_paragraph(
            "Lo ideal es conservar los nombres originales de los archivos y acompañarlos de una breve nota cuando el contexto no sea evidente. Si existen exportaciones, preferimos recibir el formato original además de PDF o captura.",
            styles["Body"],
        )
    )
    story.append(
        pdf_table(
            [
                ["Tipo de material", "Formato útil"],
                ["Expedientes de pedido", "Carpeta por pedido con presupuesto, pedido, producción, albarán, factura y comunicaciones asociadas."],
                ["Partes de máquina", "Archivo original exportado, Excel/CSV si existe, PDF o capturas si no hay exportación."],
                ["Sistemas actuales", "Capturas de pantalla con contexto: módulo, vista, filtros usados y ejemplo concreto."],
                ["Documentos internos", "Excel, Word, PDF, fotos o capturas tal como se usan actualmente."],
            ],
            [4.8 * cm, 12.6 * cm],
            styles,
        )
    )

    story.append(pdf_paragraph("5. Primer envío recomendado", styles["H1"]))
    story.append(
        pdf_paragraph(
            "Para empezar sin bloquear el trabajo, proponemos un primer paquete acotado con los siguientes elementos:",
            styles["Body"],
        )
    )
    first_pack = [
        "5-10 pedidos reales completos, incluyendo algún caso sencillo y algún caso problemático.",
        "Exportaciones o partes reales de telar, reforzadora, pulidora, disco puente y taller, si están disponibles.",
        "Capturas de las pantallas principales del sistema actual: pedido, producto, lote, stock, parte de producción, cliente y factura.",
        "Ejemplos de etiquetas físicas y trazabilidad de bloque, tabla, losa, palet, cajón y ubicación.",
        "Una breve lista de responsables por área y de las decisiones que toma cada uno.",
    ]
    story.append(
        ListFlowable(
            [ListItem(pdf_paragraph(item, styles["Bullet"])) for item in first_pack],
            bulletType="bullet",
            start="circle",
            leftIndent=12,
            bulletFontName=font,
        )
    )
    add_pdf_callout(
        story,
        "Criterio práctico:",
        "es preferible recibir documentación real aunque esté incompleta antes que esperar a preparar un paquete perfecto. Los huecos, contradicciones y formatos manuales también ayudan a detectar dónde puede aportar valor la automatización.",
        styles,
    )

    story.append(pdf_paragraph("Cierre", styles["H1"]))
    story.extend(
        [
            pdf_paragraph(
                "Con esta documentación podremos construir un mapa operativo fiable de Pulycort, validar el modelo de datos para Odoo, identificar puntos críticos de trazabilidad y priorizar pequeñas soluciones prácticas antes de abordar integraciones mayores.",
                styles["Body"],
            ),
            pdf_paragraph(
                "El primer objetivo no es auditar ni juzgar la forma actual de trabajo, sino entenderla con precisión para diseñar mejoras que encajen con la realidad de la empresa.",
                styles["Body"],
            ),
        ]
    )

    def draw_footer(canvas, pdf_doc):
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.setFillColor(hc(MUTED))
        canvas.drawString(1.6 * cm, 1.0 * cm, "Pulycort | Solicitud de documentación")
        canvas.drawRightString(A4[0] - 1.6 * cm, 1.0 * cm, f"Página {pdf_doc.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(DOCX_PATH)
    build_pdf()
    print(DOCX_PATH)
    print(PDF_PATH)


if __name__ == "__main__":
    main()
