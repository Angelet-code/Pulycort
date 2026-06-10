from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "04_analisis_y_entregables" / "informes"
DOCX_PATH = OUT_DIR / "resumen_ejecutivo_ia_pulycort.docx"


BLUE = "1F4D78"
MID_BLUE = "2E74B5"
DARK = "172033"
MUTED = "5A6678"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
WHITE = "FFFFFF"
BORDER = "D7DEE8"
RISK = "9B1C1C"
GOLD = "7A5A00"
GREEN = "1F6B4D"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn("w:{}".format(m)))
        if node is None:
            node = OxmlElement("w:{}".format(m))
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_run(run, size=None, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if size:
        run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.10):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_para(doc, text="", bold_prefix=None, color=DARK, after=6, before=0):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=before, after=after, line=1.10)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        style_run(r, size=11, bold=True, color=color)
        rest = text[len(bold_prefix):]
        if rest:
            r2 = p.add_run(rest)
            style_run(r2, size=11, color=color)
    else:
        r = p.add_run(text)
        style_run(r, size=11, color=color)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(p, after=4, line=1.167)
    r = p.add_run(text)
    style_run(r, size=10.5, color=DARK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_paragraph_spacing(p, after=4, line=1.167)
    r = p.add_run(text)
    style_run(r, size=10.5, color=DARK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style="Heading {}".format(level))
    r = p.add_run(text)
    if level == 1:
        style_run(r, size=16, bold=True, color=MID_BLUE)
        set_paragraph_spacing(p, before=16, after=8, line=1.10)
    elif level == 2:
        style_run(r, size=13, bold=True, color=MID_BLUE)
        set_paragraph_spacing(p, before=12, after=6, line=1.10)
    else:
        style_run(r, size=12, bold=True, color=BLUE)
        set_paragraph_spacing(p, before=8, after=4, line=1.10)
    return p


def add_callout(doc, title, body, fill=LIGHT_GRAY, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color=BORDER, size="4")
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=3, line=1.10)
    r = p.add_run(title)
    style_run(r, size=11, bold=True, color=accent)
    p2 = cell.add_paragraph()
    set_paragraph_spacing(p2, after=0, line=1.10)
    r2 = p2.add_run(body)
    style_run(r2, size=10.5, color=DARK)
    doc.add_paragraph()


def add_matrix(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_width(table, widths)
    hdr_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        set_cell_shading(hdr_cells[idx], LIGHT_BLUE)
        p = hdr_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, after=0, line=1.10)
        r = p.add_run(header)
        style_run(r, size=9.5, bold=True, color=BLUE)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            set_paragraph_spacing(p, after=0, line=1.10)
            if idx in (1, 2):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(value)
            style_run(r, size=9.3, color=DARK)
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell)
            set_cell_margins(cell, top=90, bottom=90)
    doc.add_paragraph()
    return table


def add_label_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    set_table_width(table, [2700, 6660])
    for label, detail in rows:
        cells = table.add_row().cells
        set_cell_shading(cells[0], LIGHT_BLUE)
        p = cells[0].paragraphs[0]
        set_paragraph_spacing(p, after=0, line=1.10)
        r = p.add_run(label)
        style_run(r, size=9.5, bold=True, color=BLUE)
        p2 = cells[1].paragraphs[0]
        set_paragraph_spacing(p2, after=0, line=1.10)
        r2 = p2.add_run(detail)
        style_run(r2, size=9.5, color=DARK)
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell)
            set_cell_margins(cell, top=90, bottom=90)
    doc.add_paragraph()
    return table


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color in [
        ("Heading 1", 16, MID_BLUE),
        ("Heading 2", 13, MID_BLUE),
        ("Heading 3", 12, BLUE),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run("Pulycort | Resumen ejecutivo IA")
    style_run(r, size=9, color=MUTED)
    return doc


def add_title_page(doc):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=4, line=1.0)
    r = p.add_run("Pulycort")
    style_run(r, size=17, bold=True, color=MUTED)

    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=6, line=1.0)
    r = p.add_run("Resumen ejecutivo")
    style_run(r, size=30, bold=True, color=BLUE)

    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=18, line=1.0)
    r = p.add_run("Pre-auditoría de procesos internos e integración de IA")
    style_run(r, size=17, color=DARK)

    add_callout(
        doc,
        "Idea central",
        "La oportunidad no está en crear una IA aislada, sino en ordenar el flujo solicitud -> presupuesto -> pedido -> producción -> albarán -> factura, usando Odoo como columna vertebral y la IA como capa de asistencia.",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )

    add_label_table(
        doc,
        [
            ("Objetivo", "Reducir trabajo manual, errores, tiempos de respuesta y dependencia de conocimiento informal."),
            ("Enfoque inicial", "Automatizaciones de baja fricción: email, documentos, facturación asistida, reporting y asistente interno."),
            ("Enfoque estratégico", "Datos industriales: producción, SIMEC, trazabilidad de bloques/tablas, stock, energía, calidad y rendimiento."),
            ("Fecha", "8 de junio de 2026."),
        ],
    )

    add_heading(doc, "Decisiones que conviene tomar primero", 2)
    for item in [
        "Confirmar qué modalidad de Odoo han contratado y qué integraciones permite.",
        "Elegir un primer piloto visible en 30-60 días, sin tocar todavía sistemas industriales críticos.",
        "Definir una fuente de verdad para clientes, pedidos, stock, producción, facturas y albaranes.",
        "Presentar la IA al equipo como ayuda para quitar trabajo repetitivo, no como vigilancia.",
    ]:
        add_bullet(doc, item)

    doc.add_page_break()


def add_content(doc):
    add_heading(doc, "1. Dónde poner el foco", 1)
    add_para(
        doc,
        "El mayor retorno inicial está en administración comercial y documentación. El mayor retorno estratégico está en producción, trazabilidad, rendimiento de material y máquina. La vía prudente es empezar por tareas con mucho volumen y poca fricción, y escalar después hacia IA industrial.",
    )

    add_matrix(
        doc,
        ["Área", "Impacto", "Fricción", "Por qué importa"],
        [
            ("Entrada de solicitudes y email", "Alto", "Baja", "Ordena pedidos que llegan por varias vías y reduce tiempo de primera respuesta."),
            ("Presupuesto -> pedido -> producción -> factura", "Muy alto", "Media", "Evita doble entrada de datos y errores entre documentos."),
            ("Facturación y albaranes", "Alto", "Media", "Hoy parece haber trabajo manual claro y dependencia de una persona."),
            ("Odoo como fuente de verdad", "Muy alto", "Media", "Debe centralizar CRM, ventas, compras, inventario, producción y facturación."),
            ("SIMEC y datos de producción", "Alto", "Media-alta", "Permite medir m2, paradas, rendimiento, merma y productividad real."),
            ("Compras, stock y trazabilidad", "Alto", "Media", "Afecta margen, plazos, disponibilidad y calidad."),
            ("Placas solares y energía", "Medio", "Baja", "Quick win: informe automático y detección de anomalías."),
        ],
        [2450, 1100, 1150, 4660],
    )

    add_callout(
        doc,
        "Recomendación",
        "No empezar por una app grande ni por visión artificial. Primero hay que asegurar datos fiables, documentos normalizados y un Odoo bien modelado.",
        fill=LIGHT_GRAY,
        accent=GREEN,
    )

    add_heading(doc, "2. Primeros pilotos recomendados", 1)
    add_matrix(
        doc,
        ["Piloto", "Qué resuelve", "Duración", "Resultado esperado"],
        [
            ("Bandeja inteligente", "Clasifica emails/formularios, extrae datos, detecta faltantes y prepara respuesta.", "4-6 semanas", "Menos tiempo administrativo y solicitudes mejor estructuradas."),
            ("Generación documental", "Crea borradores de presupuesto, pedido, orden, albarán y factura desde datos comunes.", "4-8 semanas", "Menos errores, menos copia manual y mejor trazabilidad."),
            ("Informe solar automático", "Lee informes o API de placas solares y genera resumen de rendimiento.", "1-2 semanas", "Demostración rápida de valor con baja fricción."),
            ("Diagnóstico SIMEC", "Identifica qué datos genera la máquina y cómo capturarlos.", "2-4 semanas", "Base técnica para dashboard de producción."),
        ],
        [2200, 3180, 1250, 2730],
    )

    add_heading(doc, "3. Odoo: punto crítico", 1)
    add_para(
        doc,
        "Odoo no debe tratarse solo como CRM. Si el plan contratado lo permite, debe ser la columna vertebral del flujo operativo. Antes de construir herramientas externas hay que confirmar límites de API, módulos, personalización y entorno técnico.",
    )
    add_label_table(
        doc,
        [
            ("Pregunta 1", "¿Tienen Odoo Online Standard, Odoo Online Custom, Odoo.sh u on-premise?"),
            ("Pregunta 2", "¿Hay acceso a API externa y módulos personalizados?"),
            ("Pregunta 3", "¿Cómo se modelarán bloques, tablas, lotes, acabados, stock reservado y producción?"),
            ("Pregunta 4", "¿Cómo se generarán factura y albarán, y cómo encaja con VERI*FACTU/factura electrónica?"),
            ("Decisión", "Integrar la IA con Odoo siempre que sea viable; evitar otro silo operativo."),
        ],
    )

    add_heading(doc, "4. Datos industriales: SIMEC, producción y trazabilidad", 1)
    add_para(
        doc,
        "El potencial industrial está en conectar pedido, lote, material, máquina, merma, tiempos, incidencias y calidad. Pero este bloque requiere más cuidado técnico: primero captura y visualización; después predicción y optimización.",
    )
    for item in [
        "SIMEC: identificar modelo, software, servidor interno, logs, protocolos, credenciales y formato de datos.",
        "Producción: medir m2, tiempos, paradas, causas, turnos, rendimiento por material/acabado y retrasos.",
        "Trazabilidad: vincular proveedor, bloque, tabla, lote, pedido, cliente, fotos, defectos y reclamaciones.",
        "Calidad: registrar defectos con fotos y causa probable antes de plantear visión artificial.",
    ]:
        add_bullet(doc, item)

    add_callout(
        doc,
        "Precaución técnica",
        "No conectar maquinaria industrial directamente a internet. Si hay datos de máquina, conviene capturarlos en red local con un colector edge y enviar solo datos normalizados al sistema central.",
        fill="FFF4DD",
        accent=GOLD,
    )

    doc.add_page_break()

    add_heading(doc, "5. Roadmap propuesto", 1)
    add_matrix(
        doc,
        ["Horizonte", "Objetivo", "Entregables"],
        [
            ("0-30 días", "Auditar procesos y lanzar quick wins.", "Mapa de procesos, inventario de sistemas, 5 pedidos reales analizados, prototipo de bandeja inteligente e informe solar."),
            ("30-90 días", "Automatizar flujo documental y Odoo.", "Presupuesto -> pedido -> orden -> albarán/factura con validación humana y dashboard operativo básico."),
            ("3-6 meses", "Capturar datos industriales.", "Colector SIMEC, partes de producción normalizados, stock/lotes trazables y reportes de rendimiento."),
            ("6-12 meses", "Optimizar decisiones.", "Predicción de retrasos, mantenimiento preventivo, recomendación de compras, análisis de margen real y visión artificial si hay datos suficientes."),
        ],
        [1600, 2800, 4960],
    )

    add_heading(doc, "6. Preguntas imprescindibles para la auditoría", 1)
    for item in [
        "¿Dónde se pierde más tiempo hoy: emails, presupuestos, facturas, producción, stock o incidencias?",
        "¿Qué documentos se generan manualmente y quién los genera?",
        "¿Cuántos presupuestos, pedidos, facturas y albaranes se emiten al mes?",
        "¿Qué información se copia entre sistemas o documentos?",
        "¿El stock real coincide con el stock en sistema?",
        "¿Qué datos recibe o debería recibir el servidor de SIMEC?",
        "¿De dónde salen los informes de placas solares: PDF, email, portal o API?",
        "¿Qué persona o proceso no puede fallar porque concentra demasiado conocimiento?",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "7. Dónde no conviene empezar", 1)
    add_matrix(
        doc,
        ["Evitar al principio", "Motivo"],
        [
            ("Visión artificial de calidad", "Puede ser potente, pero necesita imágenes, etiquetas y criterios consistentes."),
            ("Mantenimiento predictivo avanzado", "Sin histórico fiable de averías y paradas, sería más promesa que herramienta."),
            ("Sistema paralelo a Odoo", "Aumentaría silos y doble entrada de datos."),
            ("Marketing como primer eje", "Útil, pero el problema principal parece operativo y documental."),
            ("Automatizar decisiones sin revisión", "Al inicio la IA debe preparar y validar, no enviar ni decidir sola."),
        ],
        [2500, 6860],
    )

    add_heading(doc, "8. Mensaje interno recomendado", 1)
    add_callout(
        doc,
        "Para explicar el proyecto al equipo",
        "Queremos entender qué tareas os hacen perder más tiempo, dónde se repiten errores y qué información os falta para trabajar más cómodos. La tecnología preparará borradores, ordenará datos y quitará trabajo repetitivo; las decisiones importantes seguirán pasando por personas.",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )

    add_heading(doc, "9. Siguiente paso", 1)
    add_para(
        doc,
        "Organizar una reunión o visita de auditoría de 2-4 horas con dirección, administración, ventas, producción, compras/almacén e IT/mantenimiento.",
    )
    for item in [
        "Ver 5 pedidos reales completos de punta a punta.",
        "Confirmar plan y alcance real de Odoo.",
        "Localizar servidor SIMEC y plataforma solar.",
        "Elegir un piloto de 30-60 días.",
        "Definir métricas: horas ahorradas, errores evitados, tiempo de respuesta y visibilidad operativa.",
    ]:
        add_number(doc, item)

    add_heading(doc, "Fuentes base", 2)
    for item in [
        "Web y proceso de Pulycort: pulycort.com",
        "Documentación oficial de Odoo sobre API externa y Odoo.sh.",
        "Información oficial de AEAT sobre VERI*FACTU y facturación electrónica.",
        "Información pública de SIMEC como fabricante de maquinaria para piedra natural.",
    ]:
        add_bullet(doc, item)


def main():
    doc = setup_document()
    add_title_page(doc)
    add_content(doc)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
